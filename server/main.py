import json
from datetime import datetime
from pathlib import Path
from threading import Lock
from typing import Any

from fastapi import Body, FastAPI, Header, HTTPException
from fastapi.responses import FileResponse, JSONResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

app = FastAPI()

BASE_DIR = Path(__file__).parent
PROJECT_DIR = BASE_DIR.parent
PROFILE_PATH = BASE_DIR / "profile.json"
CONTACTS_PATH = PROJECT_DIR / "jobright_contacts.xlsx"
CONTACTS_LOCK = Lock()
CONTACT_HEADERS = [
    "Name",
    "Email",
    "Contact title",
    "Company",
    "Job title",
    "Job URL",
    "Source",
    "Captured at",
]

PROFILE_FIELDS = [
    {"key": "full_name", "label": "Full legal name", "section": "Contact", "type": "text"},
    {"key": "first_name", "label": "First name", "section": "Contact", "type": "text"},
    {"key": "middle_name", "label": "Middle name", "section": "Contact", "type": "text"},
    {"key": "last_name", "label": "Last name", "section": "Contact", "type": "text"},
    {"key": "preferred_name", "label": "Preferred name", "section": "Contact", "type": "text"},
    {"key": "pronouns", "label": "Pronouns", "section": "Contact", "type": "text"},
    {"key": "email", "label": "Email address", "section": "Contact", "type": "email"},
    {"key": "phone", "label": "Phone number", "section": "Contact", "type": "tel"},
    {"key": "address_line1", "label": "Street address", "section": "Address", "type": "text"},
    {"key": "address_line2", "label": "Apartment, suite, or unit", "section": "Address", "type": "text"},
    {"key": "city", "label": "City", "section": "Address", "type": "text"},
    {"key": "state", "label": "State or province", "section": "Address", "type": "text"},
    {"key": "postal_code", "label": "ZIP or postal code", "section": "Address", "type": "text"},
    {"key": "country", "label": "Country", "section": "Address", "type": "text"},
    {"key": "county", "label": "County", "section": "Address", "type": "text"},
    {"key": "current_title", "label": "Current job title", "section": "Experience", "type": "text"},
    {"key": "current_company", "label": "Current employer", "section": "Experience", "type": "text"},
    {"key": "years_experience", "label": "Years of experience", "section": "Experience", "type": "number"},
    {"key": "bio", "label": "Professional summary", "section": "Experience", "type": "textarea"},
    {"key": "highest_degree", "label": "Highest degree", "section": "Education", "type": "text"},
    {"key": "school", "label": "School or university", "section": "Education", "type": "text"},
    {"key": "field_of_study", "label": "Field of study", "section": "Education", "type": "text"},
    {"key": "graduation_year", "label": "Graduation year", "section": "Education", "type": "number"},
    {"key": "linkedin", "label": "LinkedIn profile URL", "section": "Links", "type": "url"},
    {"key": "portfolio", "label": "Portfolio website", "section": "Links", "type": "url"},
    {"key": "github", "label": "GitHub profile URL", "section": "Links", "type": "url"},
    {"key": "referral_source", "label": "How did you hear about this job?", "section": "Application", "type": "text"},
    {"key": "previously_worked_here", "label": "Previously worked for this organization", "section": "Application", "type": "choice", "options": ["", "Yes", "No"]},
    {"key": "phone_type", "label": "Phone device type", "section": "Contact", "type": "choice", "options": ["", "Mobile", "Landline"]},
    {"key": "phone_country_code", "label": "Phone country or calling code", "section": "Contact", "type": "text"},
    {"key": "phone_extension", "label": "Phone extension", "section": "Contact", "type": "text"},
    {"key": "authorized_to_work", "label": "Legally authorized to work", "section": "Work preferences", "type": "choice", "options": ["", "Yes", "No"]},
    {"key": "requires_sponsorship", "label": "Requires employment sponsorship", "section": "Work preferences", "type": "choice", "options": ["", "Yes", "No"]},
    {"key": "visa_status", "label": "Visa or work authorization status", "section": "Work preferences", "type": "text"},
    {"key": "work_mode", "label": "Preferred work arrangement", "section": "Work preferences", "type": "choice", "options": ["", "Remote", "Hybrid", "On-site"]},
    {"key": "employment_type", "label": "Preferred employment type", "section": "Work preferences", "type": "text"},
    {"key": "available_start_date", "label": "Available start date", "section": "Work preferences", "type": "date"},
    {"key": "notice_period", "label": "Notice period", "section": "Work preferences", "type": "text"},
    {"key": "salary_expectation", "label": "Salary expectation", "section": "Work preferences", "type": "text"},
    {"key": "willing_to_relocate", "label": "Willing to relocate", "section": "Work preferences", "type": "choice", "options": ["", "Yes", "No"]},
    {"key": "travel_willingness", "label": "Travel willingness", "section": "Work preferences", "type": "text"},
    {"key": "cover_letter", "label": "Reusable cover-letter text", "section": "Application", "type": "textarea"},
    {"key": "gender", "label": "Gender (optional)", "section": "Voluntary self-identification", "type": "text"},
    {"key": "race_ethnicity", "label": "Race or ethnicity (optional)", "section": "Voluntary self-identification", "type": "text"},
    {"key": "veteran_status", "label": "Veteran status (optional)", "section": "Voluntary self-identification", "type": "text"},
    {"key": "disability_status", "label": "Disability status (optional)", "section": "Voluntary self-identification", "type": "text"},
]

ALLOWED_PROFILE_KEYS = {field["key"] for field in PROFILE_FIELDS}


def read_profile() -> dict[str, Any]:
    try:
        data = json.loads(PROFILE_PATH.read_text(encoding="utf-8"))
    except FileNotFoundError as error:
        raise HTTPException(status_code=404, detail="Local profile not found") from error
    except json.JSONDecodeError as error:
        raise HTTPException(status_code=500, detail="Local profile contains invalid JSON") from error

    if not isinstance(data, dict):
        raise HTTPException(status_code=500, detail="Local profile must be a JSON object")

    return data


def safe_excel_text(value: Any, max_length: int = 500) -> str:
    text = str(value or "").strip()[:max_length]
    if text.startswith(("=", "+", "-", "@")):
        return f"'{text}"
    return text


def normalized_contact(raw: dict[str, Any]) -> dict[str, Any] | None:
    email = safe_excel_text(raw.get("email"), 254).lower()
    if not email or "@" not in email or any(character.isspace() for character in email):
        return None
    return {
        "Name": safe_excel_text(raw.get("name"), 120),
        "Email": email,
        "Contact title": safe_excel_text(raw.get("title"), 180),
        "Company": safe_excel_text(raw.get("company"), 160),
        "Job title": safe_excel_text(raw.get("jobTitle"), 180),
        "Job URL": safe_excel_text(raw.get("jobUrl"), 500),
        "Source": safe_excel_text(raw.get("source") or "JobRight", 80),
        "Captured at": datetime.now().astimezone().replace(tzinfo=None),
    }


def open_contacts_workbook():
    if CONTACTS_PATH.exists():
        workbook = load_workbook(CONTACTS_PATH)
        worksheet = workbook["Contacts"] if "Contacts" in workbook.sheetnames else workbook.active
        return workbook, worksheet

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Contacts"
    worksheet.append(CONTACT_HEADERS)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = "A1:H1"
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    widths = {"A": 24, "B": 34, "C": 30, "D": 22, "E": 30, "F": 58, "G": 14, "H": 22}
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width
    return workbook, worksheet


def save_contacts(records: list[dict[str, Any]]) -> dict[str, Any]:
    contacts = [contact for raw in records if (contact := normalized_contact(raw))]
    if not contacts:
        raise HTTPException(status_code=400, detail="No valid email contacts supplied")

    with CONTACTS_LOCK:
        workbook, worksheet = open_contacts_workbook()
        email_rows = {
            str(worksheet.cell(row=row, column=2).value or "").strip().lower(): row
            for row in range(2, worksheet.max_row + 1)
        }
        added = 0
        updated = 0
        for contact in contacts:
            row_number = email_rows.get(contact["Email"])
            if row_number:
                for column, header in enumerate(CONTACT_HEADERS, start=1):
                    if header == "Captured at" or contact[header]:
                        worksheet.cell(row=row_number, column=column, value=contact[header])
                updated += 1
            else:
                worksheet.append([contact[header] for header in CONTACT_HEADERS])
                email_rows[contact["Email"]] = worksheet.max_row
                added += 1

        worksheet.auto_filter.ref = f"A1:H{worksheet.max_row}"
        for cell in worksheet["H"][1:]:
            cell.number_format = "mm/dd/yy h:mm AM/PM"
        temporary_path = PROJECT_DIR / ".jobright_contacts.tmp.xlsx"
        workbook.save(temporary_path)
        temporary_path.replace(CONTACTS_PATH)

    return {
        "status": "saved",
        "added": added,
        "updated": updated,
        "total": len(email_rows),
        "path": str(CONTACTS_PATH),
    }


@app.get("/health")
def health():
    return {"status": "ok"}


@app.get("/profile")
def profile():
    return JSONResponse(read_profile(), headers={"Cache-Control": "no-store"})


@app.get("/profile-schema")
def profile_schema():
    return PROFILE_FIELDS


@app.post("/outreach-contacts")
def update_outreach_contacts(records: list[dict[str, Any]] = Body(...)):
    if len(records) > 100:
        raise HTTPException(status_code=400, detail="At most 100 contacts can be saved at once")
    return save_contacts(records)


@app.put("/profile")
def update_profile(
    updates: dict[str, Any] = Body(...),
    x_local_autofill: str | None = Header(default=None),
):
    if x_local_autofill != "profile-editor":
        raise HTTPException(status_code=403, detail="Profile editor header required")

    unknown_keys = set(updates) - ALLOWED_PROFILE_KEYS
    if unknown_keys:
        raise HTTPException(status_code=400, detail="Unknown profile field")

    for value in updates.values():
        if not isinstance(value, (str, int, float, bool)) and value is not None:
            raise HTTPException(status_code=400, detail="Profile values must be simple values")

    current = read_profile()
    current.update(updates)
    temporary_path = PROFILE_PATH.with_suffix(".json.tmp")
    temporary_path.write_text(
        json.dumps(current, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    temporary_path.replace(PROFILE_PATH)
    return {"status": "saved", "field_count": len(updates)}


@app.get("/profile-editor")
def profile_editor():
    return FileResponse(BASE_DIR / "profile-editor.html")


@app.get("/test-form")
def test_form():
    return FileResponse(BASE_DIR / "test-form.html")


@app.get("/matcher-test")
def matcher_test():
    return FileResponse(BASE_DIR / "matcher-test.html")


@app.get("/content-script-test.js")
def content_script_test():
    return FileResponse(
        BASE_DIR.parent / "extension" / "content.js",
        media_type="application/javascript",
    )


@app.get("/job-board-adapters-test.js")
def job_board_adapters_test():
    return FileResponse(
        BASE_DIR.parent / "extension" / "job-board-adapters.js",
        media_type="application/javascript",
    )
